"""
Afronomics financial model: three scenarios (conservative, base, aggressive), month by month for 36 months from
October 2026, built bottom-up from the institutional price ladder (lib/billing/ratecard.ts) plus the small reader,
sponsorship and referral lines. See docs/strategy/usd-1m-plan.md.

Every input carries a label:
    VERIFIED   checked against our own systems or code on the date given (prices, starting counts)
    ESTIMATED  reasoned from cited research, not measured (buyer pool, wedge cap)
    ASSUMED    a guess written down so it can be argued with and replaced by evidence (conversion, churn, pace)

Nothing in this model is revenue. Starting point (13 Oct 2026): USD 0 collected, 0 paying customers, 0 pilots,
6 institutions contacted on 12 Oct 2026.

How revenue is built, per scenario and month:
    contacts  -> pilots   (contact-to-pilot rate, one month later)
    pilots    -> paid     (pilot-to-paid rate, after the pilot and procurement lag)
    paid      -> Analyst / Team and API / Enterprise, by an assumed mix, at list price x a realisation factor
    customers -> churn, and Analyst -> Team upgrades
    plus custom work (a share of ladder MRR), reader subscriptions, sponsorship and remittance referrals.
New contacts are capped by the buyer pool (ESTIMATED 100-180 institutions); a contact that did not pilot can be
approached again after 12 months at a lower rate. Only the aggressive scenario opens a second buyer pool (adjacent
datasets), from the month set below, because that needs new data and capital.

    python scripts/financial_model.py   # writes docs/strategy/afronomics-financial-model.xlsx and data/strategy/model.json
"""

from __future__ import annotations

import json
import sys
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "strategy" / "afronomics-financial-model.xlsx"
MODEL_JSON = ROOT / "data" / "strategy" / "model.json"  # read by the desk's command centre (app/desk/StrategyBoard.tsx)
START = "2026-10"  # month 1 of the model
MONTHS = 36
MILESTONES = [10_000, 50_000, 100_000, 250_000, 500_000, 1_000_000]
NAMES = ("Conservative", "Base", "Aggressive")
NOT_REACHED = "not within 36 months"


@dataclass(frozen=True)
class Fact:
    label: str  # VERIFIED / ESTIMATED / ASSUMED
    value: object
    basis: str


# Where we actually are. These are not inputs to tune.
STARTING_FACTS: dict[str, Fact] = {
    "collected_usd": Fact("VERIFIED", 0, "payments table, 13 Oct 2026 (one test payment only)"),
    "paying_customers": Fact("VERIFIED", 0, "deals and paid_subscriptions tables, 13 Oct 2026"),
    "pilots": Fact("VERIFIED", 0, "no trial keys issued from the desk as of 13 Oct 2026"),
    "institutions_contacted": Fact("VERIFIED", 6, "contacted on 12 Oct 2026, logged in the desk's deal log; 6 of the 25 the D-004 test needs"),
    "prospects_logged": Fact("VERIFIED", 55, "desk deal log, stage 'prospect' (competitive-landscape.md section 6)"),
    "buyer_pool_institutions": Fact("ESTIMATED", "100-180", "serious buyers of African local-debt auction data worldwide (competitive-landscape.md)"),
    "wedge_cap_usd_year": Fact("ESTIMATED", "2-3 million", "buyer pool x the ladder at list price; the ceiling of the debt wedge alone"),
    "outside_readers_week": Fact("VERIFIED", 70, "page_views, about 70 outside readers a week (Oct 2026)"),
}

# Prices from lib/billing/ratecard.ts (institutionalLadder) and the reader rate card. Same in every scenario.
PRICE_ANALYST = 250
PRICE_TEAM = 1_500
PRICE_ENTERPRISE_FLOOR = 5_000
READER_PRO_USD = round(1_500 / 129, 2)  # KES 1,500 a month at KES_PER_USD 129


@dataclass(frozen=True)
class Input:
    key: str
    label: str
    basis: str
    values: tuple  # (Conservative, Base, Aggressive)
    unit: str = ""


INPUTS: list[Input] = [
    # Prices (VERIFIED from the code; the realised price is an assumption).
    Input("price_analyst", "VERIFIED", "institutionalLadder in lib/billing/ratecard.ts: Analyst", (PRICE_ANALYST,) * 3, "USD a month"),
    Input("price_team", "VERIFIED", "institutionalLadder: Team and API", (PRICE_TEAM,) * 3, "USD a month"),
    Input("price_enterprise", "ASSUMED", "Enterprise is 'from USD 5,000' (VERIFIED floor); the average deal above it is a guess", (5_000, 6_000, 8_000), "USD a month"),
    Input("price_realisation", "ASSUMED", "share of list actually paid after first-year discounts", (0.8, 0.9, 1.0), "x list"),
    Input("price_reader_pro", "VERIFIED", "reader Pro KES 1,500 a month at KES 129 to the dollar (ratecard.ts)", (READER_PRO_USD,) * 3, "USD a month"),
    # Pipeline.
    Input("buyer_pool", "ESTIMATED", "serious buyers of the debt wedge: low, middle, high end of 100-180", (100, 140, 180), "institutions"),
    Input("contacts_month1_extra", "ASSUMED", "further first contacts in Oct 2026 beyond the 6 already made (the D-004 test wants 25 by week 2)", (4, 14, 19), "institutions"),
    Input("contacts_per_month", "ASSUMED", "first contacts a month the founder (and later a seller) can make well", (8, 15, 25), "institutions"),
    Input("contact_to_pilot", "ASSUMED", "share of contacts that take a 30-day pilot; D-004 asks for 3 in 25 (12%)", (0.06, 0.12, 0.20), "share"),
    Input("pilot_to_paid", "ASSUMED", "share of pilots that buy; D-004 asks for 1 payment from 3 used pilots", (0.25, 0.35, 0.45), "share"),
    Input("pilot_to_paid_lag", "ASSUMED", "months from pilot start to first invoice (30-day pilot plus procurement)", (4, 3, 2), "months"),
    Input("recontact_after", "ASSUMED", "months before a contact that did not pilot is approached again", (12, 12, 12), "months"),
    Input("recontact_factor", "ASSUMED", "pilot rate on a second approach, relative to the first", (0.3, 0.5, 0.5), "x"),
    # Ladder mix of new paying customers.
    Input("mix_analyst", "ASSUMED", "share of new customers on Analyst (USD 250)", (0.60, 0.45, 0.35), "share"),
    Input("mix_team", "ASSUMED", "share on Team and API (USD 1,500)", (0.35, 0.45, 0.50), "share"),
    Input("mix_enterprise", "ASSUMED", "share on Enterprise (from USD 5,000)", (0.05, 0.10, 0.15), "share"),
    Input("upgrade_analyst_to_team", "ASSUMED", "share of Analyst customers moving to Team and API in a year", (0.05, 0.10, 0.20), "a year"),
    Input("churn_yearly", "ASSUMED", "logo churn a year, all tiers (an unproven vendor)", (0.25, 0.15, 0.10), "a year"),
    # Expansion beyond the debt wedge (adjacent datasets for the same and new buyers: FX, Eurobonds, fiscal data).
    Input("expansion_pool", "ASSUMED", "extra buyers reachable once adjacent datasets exist; aggressive only (needs capital)", (0, 0, 150), "institutions"),
    Input("expansion_from_month", "ASSUMED", "month the adjacent datasets are on sale", (0, 0, 19), "month"),
    Input("custom_work_share", "ASSUMED", "custom data work and research, as a share of ladder MRR", (0.0, 0.05, 0.10), "x MRR"),
    # Readers, sponsorship and referrals (parked per D-001 and D-007; kept as small lines).
    Input("visitors_start", "VERIFIED", "about 70 outside readers a week, roughly 300 a month (page_views)", (300, 300, 300), "a month"),
    Input("visitors_growth", "ASSUMED", "monthly growth in outside readers (wedge content, weekly wrap, monitor)", (0.04, 0.07, 0.10), "a month"),
    Input("reader_paid_conversion", "ASSUMED", "paying reader Pro subscribers as a share of monthly readers (no paywall per D-007)", (0.001, 0.002, 0.003), "share"),
    Input("sponsorship_usd", "ASSUMED", "newsletter and weekly-wrap sponsorship a month, once it starts", (0, 300, 1_000), "USD a month"),
    Input("sponsorship_from_month", "ASSUMED", "month sponsorship starts (needs an audience first)", (99, 13, 10), "month"),
    Input("referral_per_1000_visitors", "ASSUMED", "remittance referral fees per 1,000 readers; the founder has not joined the programmes yet", (0.0, 0.5, 1.5), "USD"),
    # Costs (all ASSUMED).
    Input("fixed_opex_usd", "ASSUMED", "hosting, data runs, software, legal a month", (600, 1_000, 2_500), "USD a month"),
    Input("founder_paid_from_month", "ASSUMED", "the founder draws a salary from this month", (13, 10, 4), "month"),
    Input("cost_per_head_usd", "ASSUMED", "fully loaded cost a person a month (Nairobi or remote)", (3_500, 4_500, 6_000), "USD a month"),
    Input("hires_by_month36", "ASSUMED", "people hired besides the founder by month 36, evenly", (0, 2, 6), "people"),
    Input("sales_cost_per_new_customer", "ASSUMED", "travel, events and tools per new paying institution", (1_500, 2_500, 4_000), "USD"),
]


def inp(key: str, i: int):
    for x in INPUTS:
        if x.key == key:
            return x.values[i]
    raise KeyError(key)


def month_label(m: int) -> str:
    y0, m0 = map(int, START.split("-"))
    k = y0 * 12 + (m0 - 1) + (m - 1)
    return ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"][k % 12] + f" {k // 12}"


def simulate(i: int) -> list[dict]:
    g = lambda k: inp(k, i)
    churn_m = 1 - (1 - g("churn_yearly")) ** (1 / 12)
    upgrade_m = 1 - (1 - g("upgrade_analyst_to_team")) ** (1 / 12)
    real = g("price_realisation")
    contacted_first: list[float] = []  # first contacts by month (index m-1)
    pilots_new: list[float] = []
    recontacted: list[float] = [0.0] * (MONTHS + 1)  # by original contact month, how many re-approached
    analyst = team = enterprise = 0.0
    visitors = float(g("visitors_start"))
    readers_paid = 0.0
    cash = 0.0
    ever_contacted = 0.0
    customers_ever = 0.0
    rows = []
    for m in range(1, MONTHS + 1):
        pool = g("buyer_pool") + (g("expansion_pool") if g("expansion_from_month") and m >= g("expansion_from_month") else 0)
        capacity = g("contacts_per_month") + (STARTING_FACTS["institutions_contacted"].value + g("contacts_month1_extra") - g("contacts_per_month") if m == 1 else 0)
        first = max(0.0, min(capacity, pool - ever_contacted))
        ever_contacted += first
        contacted_first.append(first)
        # Spare capacity re-approaches contacts made at least `recontact_after` months ago that did not pilot.
        spare = capacity - first
        re = 0.0
        for src in range(1, m - g("recontact_after") + 1):
            if spare <= 0:
                break
            avail = contacted_first[src - 1] * (1 - g("contact_to_pilot")) - recontacted[src]
            take = max(0.0, min(spare, avail))
            recontacted[src] += take
            spare -= take
            re += take
        # Pilots start the month after contact.
        prev_first = contacted_first[m - 2] if m >= 2 else 0.0
        prev_re = rows[-1]["recontacts"] if rows else 0.0
        p = prev_first * g("contact_to_pilot") + prev_re * g("contact_to_pilot") * g("recontact_factor")
        pilots_new.append(p)
        # Paid conversions after the lag.
        lag = g("pilot_to_paid_lag")
        paid = pilots_new[m - 1 - lag] * g("pilot_to_paid") if m - 1 - lag >= 0 else 0.0
        customers_ever += paid
        # Churn, upgrades, then new customers.
        up = analyst * upgrade_m
        analyst = analyst * (1 - churn_m) - up + paid * g("mix_analyst")
        team = team * (1 - churn_m) + up + paid * g("mix_team")
        enterprise = enterprise * (1 - churn_m) + paid * g("mix_enterprise")
        a_mrr = analyst * PRICE_ANALYST * real
        t_mrr = team * PRICE_TEAM * real
        e_mrr = enterprise * g("price_enterprise") * real
        ladder = a_mrr + t_mrr + e_mrr
        custom = ladder * g("custom_work_share")
        # Readers, sponsorship, referrals.
        if m > 1:
            visitors *= 1 + g("visitors_growth")
        readers_paid = visitors * g("reader_paid_conversion")
        reader = readers_paid * g("price_reader_pro")
        sponsor = g("sponsorship_usd") if m >= g("sponsorship_from_month") else 0.0
        referral = visitors / 1000 * g("referral_per_1000_visitors")
        revenue = ladder + custom + reader + sponsor + referral
        # Costs.
        heads = g("hires_by_month36") * m / MONTHS + (1 if m >= g("founder_paid_from_month") else 0)
        team_cost = heads * g("cost_per_head_usd")
        sales_cost = paid * g("sales_cost_per_new_customer")
        delivery = revenue * 0.08  # payment fees, data delivery, inference (ASSUMED 8% of revenue)
        opex = g("fixed_opex_usd") + team_cost + sales_cost + delivery
        profit = revenue - opex
        cash += profit
        rows.append({
            "month": m, "calendar": month_label(m), "contacts_first": round(first, 1), "recontacts": round(re, 1),
            "contacted_cumulative": round(ever_contacted, 1), "pilots_new": round(p, 2), "paid_new": round(paid, 2),
            "analyst": round(analyst, 2), "team_api": round(team, 2), "enterprise": round(enterprise, 2),
            "institutions": round(analyst + team + enterprise, 1),
            "analyst_mrr": round(a_mrr), "team_api_mrr": round(t_mrr), "enterprise_mrr": round(e_mrr), "ladder_mrr": round(ladder),
            "custom_work": round(custom), "visitors": round(visitors), "reader_subscribers": round(readers_paid, 1),
            "reader_revenue": round(reader), "sponsorship": round(sponsor), "referrals": round(referral), "revenue": round(revenue),
            "team_heads": round(heads, 2), "team_cost": round(team_cost), "sales_cost": round(sales_cost), "opex": round(opex),
            "profit": round(profit), "cash_cumulative": round(cash),
        })
    return rows


def summarise(i: int, rows: list[dict]) -> dict:
    hit = {}
    for ms in MILESTONES:
        r = next((r for r in rows if r["revenue"] >= ms), None)
        hit[ms] = {"month": r["month"], "calendar": r["calendar"]} if r else None
    trough = min(rows, key=lambda r: r["cash_cumulative"])
    last = rows[-1]
    first_paid = next((r for r in rows if r["paid_new"] > 0), None)
    return {
        "scenario": NAMES[i],
        "months_to_milestone": hit,
        "revenue_month_12": rows[11]["revenue"],
        "revenue_month_24": rows[23]["revenue"],
        "revenue_month_36": last["revenue"],
        "ladder_share_month_36": round(last["ladder_mrr"] / last["revenue"], 2) if last["revenue"] else None,
        "customers_month_36": {"analyst": round(last["analyst"], 1), "team_api": round(last["team_api"], 1), "enterprise": round(last["enterprise"], 1)},
        "contacted_by_month_36": last["contacted_cumulative"],
        "first_paid_month": first_paid["calendar"] if first_paid else None,
        "capital_required_usd": -trough["cash_cumulative"] if trough["cash_cumulative"] < 0 else 0,
        "cash_trough_month": trough["month"],
    }


def milestone_text(h: dict | None) -> str:
    return f"month {h['month']} ({h['calendar']})" if h else NOT_REACHED


def write_xlsx(results: list[tuple[list[dict], dict]]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    head = PatternFill("solid", fgColor="E8ECE6")
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    lines = [
        "Afronomics financial model: 36 months from Oct 2026, built from the institutional price ladder",
        "Generated by scripts/financial_model.py. Nothing here is revenue. Starting point 13 Oct 2026: USD 0 collected, 0 customers, 0 pilots, 6 institutions contacted.",
        "Every input is labelled VERIFIED (checked in our systems or code), ESTIMATED (reasoned from research) or ASSUMED (a guess to replace with evidence).",
        "Revenue = contacts -> pilots -> paid customers on Analyst (USD 250), Team and API (USD 1,500) and Enterprise (from USD 5,000), plus custom work, reader Pro, sponsorship and referrals.",
        "The debt wedge alone is capped at about USD 2-3m a year at list (ESTIMATED): 100-180 buyers. USD 1m a month needs the expansion path in docs/strategy/usd-1m-plan.md.",
        "Sheets: Starting facts, Inputs, Milestones, one monthly sheet per scenario. Replace ASSUMED inputs with evidence as it arrives, then re-run.",
    ]
    for n, t in enumerate(lines, 1):
        ws.cell(n, 1, t).font = Font(bold=n == 1, size=13 if n == 1 else 11)
    ws.column_dimensions["A"].width = 160

    f = wb.create_sheet("Starting facts")
    f.append(["Fact", "Label", "Value", "Basis"])
    for c in f[1]:
        c.font, c.fill = bold, head
    for k, v in STARTING_FACTS.items():
        f.append([k, v.label, v.value, v.basis])
    for col, w in zip("ABCD", (28, 12, 14, 100)):
        f.column_dimensions[col].width = w

    a = wb.create_sheet("Inputs")
    a.append(["Input", "Label", *NAMES, "Unit", "Basis"])
    for c in a[1]:
        c.font, c.fill = bold, head
    for x in INPUTS:
        a.append([x.key, x.label, *x.values, x.unit, x.basis])
    for col, w in zip("ABCDEFG", (30, 12, 14, 14, 14, 14, 100)):
        a.column_dimensions[col].width = w

    sm = wb.create_sheet("Milestones")
    sm.append(["Measure", *NAMES])
    for c in sm[1]:
        c.font, c.fill = bold, head
    for ms in MILESTONES:
        sm.append([f"First month at USD {ms:,} a month"] + [milestone_text(r[1]["months_to_milestone"][ms]) for r in results])
    for k, label in [
        ("first_paid_month", "First paying customer (model)"),
        ("revenue_month_12", "Model revenue in month 12, USD a month"),
        ("revenue_month_24", "Model revenue in month 24, USD a month"),
        ("revenue_month_36", "Model revenue in month 36, USD a month"),
        ("ladder_share_month_36", "Ladder share of revenue in month 36"),
        ("contacted_by_month_36", "Institutions contacted by month 36"),
        ("capital_required_usd", "Capital needed (deepest cumulative loss), USD"),
        ("cash_trough_month", "Month of the cash trough"),
    ]:
        sm.append([label] + [r[1][k] if r[1][k] is not None else "none" for r in results])
    sm.append(["Customers in month 36 (Analyst / Team and API / Enterprise)"] + [
        "{analyst} / {team_api} / {enterprise}".format(**r[1]["customers_month_36"]) for r in results
    ])
    sm.column_dimensions["A"].width = 60
    for col in "BCD":
        sm.column_dimensions[col].width = 26

    for n, (rows, _) in enumerate(results):
        w = wb.create_sheet(NAMES[n])
        w.append(list(rows[0].keys()))
        for c in w[1]:
            c.font, c.fill = bold, head
        for r in rows:
            w.append(list(r.values()))
        w.freeze_panes = "C2"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)


def main() -> int:
    for i in range(3):
        mix = inp("mix_analyst", i) + inp("mix_team", i) + inp("mix_enterprise", i)
        assert abs(mix - 1) < 1e-9, f"ladder mix for {NAMES[i]} sums to {mix}"
    results = [(rows, summarise(i, rows)) for i, rows in ((i, simulate(i)) for i in range(3))]
    write_xlsx(results)
    MODEL_JSON.parent.mkdir(parents=True, exist_ok=True)
    MODEL_JSON.write_text(json.dumps({
        "note": "Monthly revenue (USD) per scenario from scripts/financial_model.py, built from the price ladder. ASSUMPTIONS, not revenue. Month 1 = start.",
        "start": START,
        "scenarios": {NAMES[n]: [r["revenue"] for r in rows] for n, (rows, _) in enumerate(results)},
        "institutions": {NAMES[n]: [r["institutions"] for r in rows] for n, (rows, _) in enumerate(results)},
        "ladder": {"analyst": PRICE_ANALYST, "team_api": PRICE_TEAM, "enterprise_from": PRICE_ENTERPRISE_FLOOR},
        "milestones": {NAMES[n]: {str(ms): milestone_text(s["months_to_milestone"][ms]) for ms in MILESTONES} for n, (_, s) in enumerate(results)},
        "starting_facts": {k: {"label": v.label, "value": v.value, "basis": v.basis} for k, v in STARTING_FACTS.items()},
    }, indent=1), encoding="utf-8")
    print(json.dumps([s for _, s in results], indent=1, default=str))
    print(f"wrote {OUT.relative_to(ROOT)} and {MODEL_JSON.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
